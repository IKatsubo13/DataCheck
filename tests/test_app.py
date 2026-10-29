from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
APP_PATH = ROOT / "app.py"
SAMPLE_DIR = ROOT / "sample_data"


def start_app() -> AppTest:
    app = AppTest.from_file(str(APP_PATH), default_timeout=120)
    app.run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    return app


def click_compare(app: AppTest) -> None:
    next(button for button in app.button if "Сверить файлы" in button.label).click().run()
    assert not list(app.exception), [exc.value for exc in app.exception]


def compare_files(
    file_a: bytes,
    file_b: bytes,
    name_a: str = "file_a.csv",
    name_b: str = "file_b.csv",
) -> AppTest:
    app = start_app()
    app.file_uploader[0].set_value((name_a, file_a, "text/csv"))
    app.file_uploader[1].set_value((name_b, file_b, "text/csv"))
    app.run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)
    return app


def find_diff(app: AppTest) -> pd.DataFrame:
    for element in app.dataframe:
        frame = element.value
        if {"Ключ", "Поле", "Тип", "Файл A", "Файл B"} <= set(frame.columns):
            return frame
    raise AssertionError("Таблица расхождений не найдена")


def test_app_starts_with_two_uploaders():
    app = start_app()
    assert len(app.file_uploader) == 2
    assert not app.button


def test_demo_finds_typed_differences_and_builds_excel():
    app = compare_files(
        (SAMPLE_DIR / "operational_data.csv").read_bytes(),
        (SAMPLE_DIR / "accounting_data.csv").read_bytes(),
        "operational_data.csv",
        "accounting_data.csv",
    )

    assert [tab.label for tab in app.tabs] == [
        "01 · Расхождения · 4",
        "02 · Полнота · 4",
        "03 · Ключи · 4",
        "04 · Методика",
    ]

    diff = find_diff(app)
    assert len(diff) == 4
    assert set(zip(diff["Ключ"], diff["Поле"], diff["Тип"])) == {
        ("A-002", "cost_rub", "Число"),
        ("A-003", "status", "Текст"),
        ("A-006", "owner", "Текст"),
        ("A-006", "renewal_date", "Дата"),
    }
    assert set(diff["Ключ"]) == {"A-002", "A-003", "A-006"}

    markup = "\n".join(str(element.value) for element in app.markdown)
    assert "renewal_date — дата" in markup

    result = app.session_state["dc_result"]
    workbook = load_workbook(BytesIO(result.report))
    assert workbook.sheetnames == [
        "Сводка",
        "Параметры",
        "Расхождения",
        "Только_A",
        "Только_B",
        "Дубли_A",
        "Дубли_B",
        "Без_ключа_A",
        "Без_ключа_B",
    ]
    diff_rows = list(workbook["Расхождения"].iter_rows(values_only=True))
    assert diff_rows[0] == ("Ключ", "Поле", "Тип", "Файл A", "Файл B", "Строка A", "Строка B")
    assert len(diff_rows) == 5
    parameters = dict(workbook["Параметры"].iter_rows(min_row=2, values_only=True))
    assert "renewal_date — Дата" in parameters["Типы полей"]
    assert parameters["Допуск для чисел"] == "±0.01"
    assert parameters["Относительный допуск"] == "0.00%"
    assert parameters["Пробелы по краям текста"] == "игнорируются"
    assert parameters["Регистр текста"] == "игнорируется"
    assert parameters["Пустые значения"] == "учитываются"
    assert parameters["Формат дат"] == "игнорируется"


def test_number_and_date_formats_are_normalized():
    file_a = """asset_id,amount,renewal_date,owner
A-1,"1 234,50",2024-02-29,Alice
A-2,"1,234.50",2024-03-01 10:00:00,Bob
A-3,"(1 000,00)",2023-12-31,Carol
A-4,2000,2024-01-01,Dan
"""
    file_b = """asset_id,amount,renewal_date,owner
A-1,1234.5,29.02.2024,Alice
A-2,"1234,50",01.03.2024 15:30:00,Bob
A-3,"-1000",31.12.2023,Carol
A-4,"2 000,50",02.01.2024,Dan
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    diff = find_diff(app)
    assert len(diff) == 2
    assert set(zip(diff["Ключ"], diff["Поле"], diff["Тип"])) == {
        ("A-4", "amount", "Число"),
        ("A-4", "renewal_date", "Дата"),
    }


def test_leading_zero_codes_remain_text():
    file_a = """asset_id,code,owner
ID-007,007,Alice
ID-008,42,Bob
"""
    file_b = """asset_id,code,owner
ID-007,7,Alice
ID-008,42,Bob
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    assert "02 · Полнота · 0" in [tab.label for tab in app.tabs]
    diff = find_diff(app)
    assert len(diff) == 1
    row = diff.iloc[0]
    assert (row["Ключ"], row["Поле"], row["Тип"]) == ("ID-007", "code", "Текст")
    assert (row["Файл A"], row["Файл B"]) == ("007", "7")


def test_text_edge_spaces_are_ignored_by_default_and_can_be_enabled():
    file_a = """asset_id,owner,status
A-1," Alice "," Активна "
A-2,Bob,Активна
"""
    file_b = """asset_id,owner,status
A-1,Alice,Активна
A-2,Bob,Активна
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    assert "01 · Расхождения · 0" in [tab.label for tab in app.tabs]
    spaces = next(box for box in app.checkbox if "пробелы" in box.label.lower())
    assert spaces.value is True

    spaces.set_value(False).run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)

    diff = find_diff(app)
    assert set(zip(diff["Ключ"], diff["Поле"], diff["Тип"])) == {
        ("A-1", "owner", "Текст"),
        ("A-1", "status", "Текст"),
    }


def test_date_format_rule_is_configurable():
    file_a = """asset_id,renewal_date
A-1,2024-04-12
A-2,2024-05-01
"""
    file_b = """asset_id,renewal_date
A-1,12.04.2024
A-2,2024-05-01
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    assert "01 · Расхождения · 0" in [tab.label for tab in app.tabs]
    date_rule = next(box for box in app.checkbox if "формате дат" in box.label.lower())
    assert date_rule.value is True

    date_rule.set_value(False).run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)

    diff = find_diff(app)
    assert len(diff) == 1
    row = diff.iloc[0]
    assert (row["Ключ"], row["Поле"], row["Тип"]) == ("A-1", "renewal_date", "Дата")


def test_empty_values_rule_is_configurable():
    file_a = """asset_id,owner,status
A-1,,Активна
A-2,Alice,Активна
"""
    file_b = """asset_id,owner,status
A-1,ИТ,Активна
A-2,Alice,Активна
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    diff = find_diff(app)
    assert len(diff) == 1
    assert (diff.iloc[0]["Ключ"], diff.iloc[0]["Поле"]) == ("A-1", "owner")

    empty_rule = next(box for box in app.checkbox if "пустые" in box.label.lower())
    assert empty_rule.value is False
    empty_rule.set_value(True).run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)

    assert "01 · Расхождения · 0" in [tab.label for tab in app.tabs]
    markup = "\n".join(str(element.value) for element in app.markdown)
    assert "Пустые значения игнорируются" in markup


def test_text_case_is_ignored_by_default_and_can_be_enabled():
    file_a = """asset_id,status,owner
A-1,Активна,Alice
A-2,Активна,Bob
"""
    file_b = """asset_id,status,owner
A-1,активна,ALICE
A-2,Активна,Bob
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    assert "01 · Расхождения · 0" in [tab.label for tab in app.tabs]
    case_rule = next(box for box in app.checkbox if "регистр" in box.label.lower())
    assert case_rule.value is True

    case_rule.set_value(False).run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)

    diff = find_diff(app)
    assert set(zip(diff["Ключ"], diff["Поле"], diff["Тип"])) == {
        ("A-1", "status", "Текст"),
        ("A-1", "owner", "Текст"),
    }


def test_absolute_numeric_tolerance_is_configurable():
    file_a = """asset_id,amount
A-1,1000
A-2,500
"""
    file_b = """asset_id,amount
A-1,1000.01
A-2,500.02
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    diff = find_diff(app)
    assert len(diff) == 1
    assert (diff.iloc[0]["Ключ"], diff.iloc[0]["Поле"]) == ("A-2", "amount")

    tolerance = next(widget for widget in app.number_input if widget.label == "Допуск для чисел")
    assert tolerance.value == 0.01
    tolerance.set_value(0.02).run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)

    assert "01 · Расхождения · 0" in [tab.label for tab in app.tabs]
    assert "±0.02" in "\n".join(str(element.value) for element in app.markdown)


def test_relative_numeric_tolerance_is_configurable():
    file_a = """asset_id,amount
A-1,1000
A-2,2000
"""
    file_b = """asset_id,amount
A-1,1004
A-2,2012
"""
    app = compare_files(file_a.encode("utf-8"), file_b.encode("utf-8"))

    assert "01 · Расхождения · 2" in [tab.label for tab in app.tabs]

    relative = next(
        widget for widget in app.number_input if widget.label == "Относительный допуск, %"
    )
    assert relative.value == 0.0
    relative.set_value(0.5).run()
    assert not list(app.exception), [exc.value for exc in app.exception]
    click_compare(app)

    diff = find_diff(app)
    assert len(diff) == 1
    assert diff.iloc[0]["Ключ"] == "A-2"
    assert "01 · Расхождения · 1" in [tab.label for tab in app.tabs]

    relative = next(
        widget for widget in app.number_input if widget.label == "Относительный допуск, %"
    )
    relative.set_value(0.6).run()
    click_compare(app)
    assert "01 · Расхождения · 0" in [tab.label for tab in app.tabs]
