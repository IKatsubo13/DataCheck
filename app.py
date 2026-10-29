from __future__ import annotations

import hashlib
import html
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FAVICON = Path(__file__).resolve().parent / "assets" / "favicon.png"

st.set_page_config(
    page_title="DataCheck — аудит CSV",
    page_icon=str(FAVICON) if FAVICON.exists() else None,
    layout="wide",
    initial_sidebar_state="collapsed",
)


CSS = """
@import url('https://fonts.googleapis.com/css2?family=Fraunces:ital,opsz,wght@0,9..144,400;0,9..144,500;0,9..144,600;1,9..144,400;1,9..144,500;1,9..144,600&family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Sans:wght@400;500;600&display=swap');

:root {
    --paper: #f3eee3;
    --paper-high: #fcf8f0;
    --ink: #14221d;
    --ink-soft: #53625b;
    --line: rgba(20, 34, 29, 0.15);
    --line-strong: rgba(20, 34, 29, 0.38);
    --accent: #c34a2c;
    --accent-soft: #f1d9ce;
    --green: #2f6a55;
    --green-soft: #ddeae0;
    --amber: #9a6414;
    --amber-soft: #f4e6c8;
    --red: #a62f2f;
    --red-soft: #f4ddd7;
    --shadow: 0 24px 60px rgba(20, 34, 29, 0.09);
    --radius: 18px;
}

html,
body,
[class*="css"] {
    font-family: "IBM Plex Sans", "Segoe UI", sans-serif;
    color: var(--ink);
}

.stApp {
    background-color: var(--paper);
    background-image:
        radial-gradient(1100px 520px at 88% -8%, rgba(195, 74, 44, 0.13), transparent 62%),
        radial-gradient(900px 500px at -8% 12%, rgba(47, 106, 85, 0.10), transparent 58%),
        linear-gradient(rgba(20, 34, 29, 0.035) 1px, transparent 1px),
        linear-gradient(90deg, rgba(20, 34, 29, 0.035) 1px, transparent 1px);
    background-size: auto, auto, 30px 30px, 30px 30px;
}

.stApp::before {
    content: "";
    position: fixed;
    inset: 0 0 auto 0;
    height: 5px;
    background: var(--accent);
    z-index: 9999;
}

[data-testid="stHeader"] {
    background: transparent;
}

[data-testid="stAppDeployButton"],
#MainMenu,
footer:not(.dc-footer) {
    display: none !important;
}

.block-container {
    max-width: 1240px;
    padding-top: 2.8rem;
    padding-bottom: 2.5rem;
}

h1,
h2,
h3,
h4 {
    font-family: "Fraunces", Georgia, "Times New Roman", serif !important;
    color: var(--ink);
    font-weight: 500;
    letter-spacing: -0.025em;
}

p,
li,
label {
    font-family: "IBM Plex Sans", "Segoe UI", sans-serif;
}

code {
    font-family: "IBM Plex Mono", Consolas, monospace;
    background: rgba(20, 34, 29, 0.07);
    border-radius: 5px;
    padding: 0.1rem 0.3rem;
}

[data-testid="stCaptionContainer"] p {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.68rem;
    letter-spacing: 0.02em;
}

/* Hero */
.dc-hero {
    margin-bottom: 2.4rem;
}

.dc-brandline {
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 1rem;
    margin-bottom: 2.2rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.68rem;
    letter-spacing: 0.08em;
    text-transform: uppercase;
}

.dc-brand {
    display: flex;
    align-items: center;
    gap: 0.6rem;
    color: var(--ink);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.72rem;
    font-weight: 500;
    letter-spacing: 0.14em;
}

.dc-brand-mark {
    position: relative;
    display: inline-block;
    width: 18px;
    height: 18px;
    border: 1.5px solid var(--ink);
    border-radius: 50%;
}

.dc-brand-mark::after {
    content: "";
    position: absolute;
    inset: 3px;
    border-radius: 50%;
    background: var(--accent);
}

.dc-hero-grid {
    display: grid;
    grid-template-columns: minmax(0, 1.2fr) minmax(300px, 0.8fr);
    gap: clamp(2rem, 5vw, 5rem);
    align-items: center;
}

.dc-hero-copy {
    animation: dc-rise 650ms cubic-bezier(0.22, 1, 0.36, 1) both;
}

.dc-eyebrow {
    display: flex;
    align-items: center;
    gap: 0.75rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.66rem;
    letter-spacing: 0.13em;
    text-transform: uppercase;
}

.dc-eyebrow::before {
    content: "";
    display: inline-block;
    width: 28px;
    height: 1px;
    background: var(--accent);
}

.dc-version {
    border: 1px solid var(--line);
    border-radius: 999px;
    padding: 0.22rem 0.48rem;
    color: var(--accent);
    font-size: 0.6rem;
}

.dc-hero h1 {
    max-width: 760px;
    margin: 1rem 0 1.15rem;
    font-size: clamp(3.3rem, 6.7vw, 6.1rem);
    line-height: 0.92;
    font-weight: 500;
}

.dc-hero h1 em {
    color: var(--accent);
    font-weight: 400;
}

.dc-lede {
    max-width: 41rem;
    margin: 0;
    color: var(--ink-soft);
    font-size: 1.06rem;
    line-height: 1.68;
}

.dc-meta {
    display: flex;
    flex-wrap: wrap;
    gap: 0.5rem;
    margin-top: 1.6rem;
}

.dc-meta span {
    padding: 0.35rem 0.6rem;
    border: 1px solid var(--line);
    border-radius: 999px;
    background: rgba(252, 248, 240, 0.52);
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.63rem;
    letter-spacing: 0.05em;
    text-transform: uppercase;
}

.dc-art {
    position: relative;
    min-height: 336px;
    padding: 1.35rem;
    overflow: hidden;
    border: 1px solid var(--line);
    border-radius: 26px;
    background:
        linear-gradient(145deg, rgba(252, 248, 240, 0.97), rgba(243, 238, 227, 0.76));
    box-shadow: var(--shadow);
    transform: rotate(1.2deg);
    animation: dc-rise 800ms 100ms cubic-bezier(0.22, 1, 0.36, 1) both;
}

.dc-art::before {
    content: "";
    position: absolute;
    inset: 0;
    pointer-events: none;
    background-image: radial-gradient(rgba(20, 34, 29, 0.22) 0.7px, transparent 0.7px);
    background-size: 15px 15px;
    opacity: 0.3;
}

.dc-art::after {
    content: "";
    position: absolute;
    right: -70px;
    bottom: -110px;
    width: 240px;
    height: 240px;
    border: 1px solid rgba(195, 74, 44, 0.3);
    border-radius: 50%;
}

.dc-art-top,
.dc-art-foot {
    position: relative;
    z-index: 1;
    display: flex;
    justify-content: space-between;
    gap: 1rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.6rem;
    letter-spacing: 0.1em;
    text-transform: uppercase;
}

.dc-filepair {
    position: relative;
    z-index: 1;
    display: grid;
    grid-template-columns: 1fr auto 1fr;
    gap: 0.65rem;
    align-items: center;
    margin-top: 1.3rem;
}

.dc-mini-file {
    min-width: 0;
    padding: 0.85rem 0.8rem;
    border: 1px solid var(--line);
    border-radius: 12px;
    background: var(--paper-high);
    box-shadow: 0 10px 24px rgba(20, 34, 29, 0.06);
}

.dc-mini-file--b {
    transform: translateY(18px);
}

.dc-mini-tag {
    display: block;
    margin-bottom: 0.5rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.58rem;
    letter-spacing: 0.08em;
}

.dc-mini-row {
    display: flex;
    justify-content: space-between;
    gap: 0.5rem;
    padding: 0.32rem 0;
    border-bottom: 1px dashed var(--line);
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.59rem;
}

.dc-mini-row:last-child {
    border-bottom: none;
}

.dc-mini-row strong {
    color: var(--ink);
    font-weight: 500;
}

.dc-mini-row .dc-alert {
    color: var(--accent);
}

.dc-bridge {
    display: grid;
    place-items: center;
    width: 34px;
    height: 34px;
    border-radius: 50%;
    background: var(--accent);
    color: #fff7ec;
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.9rem;
    box-shadow: 0 0 0 7px var(--paper);
}

.dc-art-foot {
    margin-top: 2rem;
}

/* Sections */
.dc-section {
    display: flex;
    align-items: baseline;
    gap: 0.85rem;
    margin: 2.8rem 0 1rem;
    padding-top: 0.95rem;
    border-top: 1px solid var(--line);
}

.dc-section-index {
    color: var(--accent);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.66rem;
    letter-spacing: 0.05em;
}

.dc-section-title {
    font-family: "Fraunces", Georgia, serif;
    font-size: 1.45rem;
    letter-spacing: -0.02em;
}

.dc-section-note {
    margin-left: auto;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.6rem;
    letter-spacing: 0.04em;
    text-transform: uppercase;
}

/* Upload */
.dc-upload-head {
    display: flex;
    gap: 0.7rem;
    min-height: 86px;
    margin-bottom: 0.45rem;
}

.dc-upload-tag {
    display: grid;
    place-items: center;
    flex: 0 0 auto;
    width: 28px;
    height: 28px;
    border-radius: 7px;
    background: var(--ink);
    color: var(--paper-high);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.7rem;
}

.dc-upload-title {
    margin-top: 0.1rem;
    font-family: "Fraunces", Georgia, serif;
    font-size: 1.12rem;
    font-weight: 500;
}

.dc-upload-copy {
    max-width: 32rem;
    margin: 0.25rem 0 0;
    color: var(--ink-soft);
    font-size: 0.82rem;
    line-height: 1.5;
}

[data-testid="stFileUploader"] {
    margin-top: 0;
}

[data-testid="stFileUploaderDropzone"] {
    min-height: 124px !important;
    padding: 1rem !important;
    border: 1px dashed var(--line-strong) !important;
    border-radius: 14px !important;
    background: rgba(252, 248, 240, 0.7) !important;
    transition: border-color 180ms ease, background 180ms ease;
}

[data-testid="stFileUploaderDropzone"]:hover {
    border-color: var(--accent) !important;
    background: var(--paper-high) !important;
}

[data-testid="stFileUploaderDropzone"] button {
    border: 1px solid var(--ink) !important;
    border-radius: 999px !important;
    background: var(--ink) !important;
    color: var(--paper-high) !important;
    font-family: "IBM Plex Mono", Consolas, monospace !important;
    font-size: 0.64rem !important;
    letter-spacing: 0.03em !important;
}

[data-testid="stFileUploaderDropzone"] button:hover {
    border-color: var(--accent) !important;
    background: var(--accent) !important;
}

[data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-borderlessIcon"],
[data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-minimal"] {
    border: none !important;
    background: transparent !important;
    color: var(--ink-soft) !important;
}

[data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-borderlessIcon"]:hover,
[data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-minimal"]:hover {
    background: rgba(20, 34, 29, 0.08) !important;
    color: var(--ink) !important;
}

[data-testid="stFileUploaderDropzone"] small,
[data-testid="stFileUploaderDropzone"] span {
    color: var(--ink-soft);
}

[data-testid="stFileUploaderDropzone"] svg {
    color: var(--ink-soft);
}

/* Guide */
.dc-guide {
    display: grid;
    grid-template-columns: repeat(3, minmax(0, 1fr));
    gap: 1px;
    margin-top: 1.5rem;
    overflow: hidden;
    border: 1px solid var(--line);
    border-radius: var(--radius);
    background: var(--line);
}

.dc-guide-card {
    min-height: 154px;
    padding: 1.3rem 1.35rem;
    background: rgba(252, 248, 240, 0.74);
}

.dc-guide-num {
    color: var(--accent);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.62rem;
    letter-spacing: 0.08em;
}

.dc-guide-card h3 {
    margin: 1.1rem 0 0.4rem;
    font-size: 1.18rem;
    line-height: 1.1;
}

.dc-guide-card p {
    margin: 0;
    color: var(--ink-soft);
    font-size: 0.84rem;
    line-height: 1.55;
}

.dc-demo-hint {
    margin-top: 0.8rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.66rem;
    line-height: 1.7;
}

.dc-demo-hint strong {
    color: var(--accent);
    font-weight: 500;
}

/* Settings */
[data-testid="stVerticalBlockBorderWrapper"] {
    border: 1px solid var(--line) !important;
    border-radius: var(--radius) !important;
    background: rgba(252, 248, 240, 0.68);
    box-shadow: 0 12px 36px rgba(20, 34, 29, 0.04);
}

[data-testid="stSelectbox"] label p,
[data-testid="stMultiSelect"] label p,
[data-testid="stNumberInput"] label p {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.63rem;
    letter-spacing: 0.06em;
    text-transform: uppercase;
}

[data-testid="stNumberInput"] [data-baseweb="input"] {
    border-color: var(--line) !important;
    border-radius: 11px !important;
    background: var(--paper-high) !important;
}

[data-testid="stNumberInput"] input {
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.8rem;
}

[data-testid="stCheckbox"] label p {
    color: var(--ink);
    font-size: 0.84rem;
}

[data-baseweb="select"] > div {
    border-color: var(--line) !important;
    border-radius: 11px !important;
    background: var(--paper-high) !important;
}

.dc-inline-note {
    padding-top: 0.8rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.65rem;
    line-height: 1.55;
}

.dc-tolerance-note {
    padding-bottom: 0.2rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.62rem;
    line-height: 1.6;
}

.dc-change-note {
    margin-top: 0.65rem;
    padding: 0.85rem 1rem;
    border: 1px dashed var(--line-strong);
    border-radius: 12px;
    background: rgba(244, 230, 200, 0.38);
    color: var(--ink-soft);
    font-size: 0.84rem;
}

.stButton > button,
.stDownloadButton > button,
[data-testid="stBaseButton-primary"],
[data-testid="stBaseButton-secondary"] {
    font-family: "IBM Plex Mono", Consolas, monospace !important;
    font-size: 0.68rem !important;
    letter-spacing: 0.03em !important;
    border-radius: 999px !important;
    transition: transform 160ms ease, background 160ms ease, border-color 160ms ease, box-shadow 160ms ease !important;
}

[data-testid="stBaseButton-primary"] {
    border: 1px solid var(--ink) !important;
    background: var(--ink) !important;
    color: var(--paper-high) !important;
    box-shadow: 0 8px 20px rgba(20, 34, 29, 0.13);
}

[data-testid="stBaseButton-primary"]:hover {
    border-color: var(--accent) !important;
    background: var(--accent) !important;
    color: #fff7ec !important;
    transform: translateY(-1px);
}

[data-testid="stBaseButton-secondary"] {
    border: 1px solid var(--line-strong) !important;
    background: var(--paper-high) !important;
    color: var(--ink) !important;
}

[data-testid="stBaseButton-secondary"]:hover {
    border-color: var(--accent) !important;
    color: var(--accent) !important;
    transform: translateY(-1px);
}

/* Expanders and tables */
[data-testid="stExpander"] details {
    border: 1px solid var(--line) !important;
    border-radius: 13px !important;
    background: rgba(252, 248, 240, 0.58);
}

[data-testid="stExpander"] summary {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.67rem;
}

[data-testid="stDataFrame"] {
    overflow: hidden;
    border: 1px solid var(--line);
    border-radius: 13px;
    box-shadow: 0 6px 24px rgba(20, 34, 29, 0.035);
}

/* Results */
.dc-runline {
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    gap: 0.5rem;
    margin: -0.25rem 0 1rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.64rem;
}

.dc-runline-label {
    letter-spacing: 0.08em;
    text-transform: uppercase;
}

.dc-runline-file {
    max-width: 280px;
    overflow: hidden;
    color: var(--ink);
    text-overflow: ellipsis;
    white-space: nowrap;
}

.dc-runline-x {
    color: var(--accent);
}

.dc-runline-divider {
    width: 1px;
    height: 12px;
    margin: 0 0.25rem;
    background: var(--line-strong);
}

.dc-verdict {
    display: grid;
    grid-template-columns: minmax(0, 1fr) auto;
    gap: 2rem;
    align-items: center;
    margin-bottom: 0.9rem;
    padding: 1.6rem 1.8rem;
    border: 1px solid var(--line);
    border-left: 4px solid var(--amber);
    border-radius: var(--radius);
    background: linear-gradient(135deg, rgba(252, 248, 240, 0.96), rgba(244, 230, 200, 0.45));
    box-shadow: var(--shadow);
    animation: dc-rise 500ms cubic-bezier(0.22, 1, 0.36, 1) both;
}

.dc-verdict--critical {
    border-left-color: var(--red);
    background: linear-gradient(135deg, rgba(252, 248, 240, 0.96), rgba(244, 221, 215, 0.58));
}

.dc-verdict--clear {
    border-left-color: var(--green);
    background: linear-gradient(135deg, rgba(252, 248, 240, 0.96), rgba(221, 234, 224, 0.62));
}

.dc-verdict-kicker {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.62rem;
    letter-spacing: 0.11em;
    text-transform: uppercase;
}

.dc-verdict h2 {
    margin: 0.5rem 0 0.65rem;
    font-size: clamp(1.8rem, 3.4vw, 2.85rem);
    line-height: 1.02;
    font-weight: 500;
}

.dc-verdict p {
    max-width: 58ch;
    margin: 0;
    color: var(--ink-soft);
    font-size: 0.9rem;
    line-height: 1.62;
}

.dc-stamp {
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
    width: 122px;
    height: 122px;
    border: 1.5px solid currentColor;
    border-radius: 50%;
    outline: 1px solid currentColor;
    outline-offset: 5px;
    color: var(--amber);
    background: rgba(252, 248, 240, 0.42);
    transform: rotate(-7deg);
    font-family: "IBM Plex Mono", Consolas, monospace;
    text-align: center;
    text-transform: uppercase;
}

.dc-verdict--critical .dc-stamp {
    color: var(--red);
}

.dc-verdict--clear .dc-stamp {
    color: var(--green);
}

.dc-stamp span {
    font-size: 0.53rem;
    letter-spacing: 0.12em;
}

.dc-stamp strong {
    margin: 0.25rem 0;
    font-size: 0.8rem;
    font-weight: 500;
    letter-spacing: 0.08em;
}

.dc-stats {
    display: grid;
    grid-template-columns: repeat(5, minmax(0, 1fr));
    gap: 0.7rem;
    margin: 0.9rem 0;
}

.dc-stat {
    position: relative;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    min-height: 126px;
    padding: 0.95rem 1rem 0.85rem;
    overflow: hidden;
    border: 1px solid var(--line);
    border-radius: 14px;
    background: rgba(252, 248, 240, 0.82);
    animation: dc-rise 500ms cubic-bezier(0.22, 1, 0.36, 1) both;
}

.dc-stat:nth-child(2) { animation-delay: 50ms; }
.dc-stat:nth-child(3) { animation-delay: 100ms; }
.dc-stat:nth-child(4) { animation-delay: 150ms; }
.dc-stat:nth-child(5) { animation-delay: 200ms; }

.dc-stat::before {
    content: "";
    position: absolute;
    inset: 0 0 auto;
    height: 3px;
    background: var(--line-strong);
}

.dc-stat--critical::before { background: var(--red); }
.dc-stat--warning::before { background: var(--amber); }
.dc-stat--clear::before { background: var(--green); }

.dc-stat-index {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.57rem;
    letter-spacing: 0.06em;
}

.dc-stat-label {
    margin-top: 0.45rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.61rem;
    letter-spacing: 0.03em;
    line-height: 1.4;
    text-transform: uppercase;
}

.dc-stat-value {
    margin: 0.35rem 0 0.3rem;
    color: var(--ink);
    font-family: "Fraunces", Georgia, serif;
    font-size: 2.2rem;
    line-height: 1;
    letter-spacing: -0.04em;
}

.dc-stat-value .dc-slash {
    margin: 0 0.18rem;
    color: var(--line-strong);
    font-size: 1.35rem;
}

.dc-stat-foot {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.58rem;
    line-height: 1.45;
}

.dc-priority {
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    gap: 0.45rem;
    margin: 0.35rem 0 1.15rem;
}

.dc-priority-label {
    margin-right: 0.35rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.6rem;
    letter-spacing: 0.07em;
    text-transform: uppercase;
}

.dc-chip {
    display: inline-flex;
    align-items: center;
    gap: 0.45rem;
    padding: 0.36rem 0.64rem;
    border: 1px solid var(--line);
    border-radius: 999px;
    background: rgba(252, 248, 240, 0.7);
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.62rem;
    letter-spacing: 0.02em;
}

.dc-chip i {
    color: var(--ink-soft);
    font-style: normal;
    font-size: 0.55rem;
}

.dc-chip b {
    color: var(--ink);
    font-weight: 500;
}

.dc-chip--critical {
    border-color: rgba(166, 47, 47, 0.3);
    background: var(--red-soft);
    color: var(--red);
}

.dc-chip--critical i,
.dc-chip--critical b {
    color: var(--red);
}

.dc-chip--warning {
    border-color: rgba(154, 100, 20, 0.3);
    background: var(--amber-soft);
    color: var(--amber);
}

.dc-chip--warning i,
.dc-chip--warning b {
    color: var(--amber);
}

.dc-chip--clear {
    border-color: rgba(47, 106, 85, 0.25);
    background: rgba(221, 234, 224, 0.7);
    color: var(--green);
}

.dc-chip--clear i,
.dc-chip--clear b {
    color: var(--green);
}

.dc-action-note {
    padding-top: 0.7rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.64rem;
    line-height: 1.55;
}

.dc-table-head {
    display: flex;
    align-items: baseline;
    justify-content: space-between;
    gap: 1rem;
    margin: 0.25rem 0 0.6rem;
}

.dc-table-title {
    font-family: "Fraunces", Georgia, serif;
    font-size: 1.18rem;
    letter-spacing: -0.015em;
}

.dc-table-meta {
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.6rem;
    letter-spacing: 0.03em;
    text-transform: uppercase;
}

.dc-tab-lede {
    max-width: 66ch;
    margin: 0 0 1rem;
    color: var(--ink-soft);
    font-size: 0.86rem;
    line-height: 1.6;
}

.dc-table-note {
    margin-top: 0.55rem;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.6rem;
    line-height: 1.6;
}

.dc-note-box {
    margin-bottom: 1rem;
    padding: 0.9rem 1rem;
    border: 1px dashed var(--line-strong);
    border-radius: 13px;
    background: rgba(252, 248, 240, 0.58);
    color: var(--ink-soft);
    font-size: 0.84rem;
    line-height: 1.6;
}

.dc-note-box strong {
    color: var(--ink);
    font-weight: 500;
}

.dc-clear {
    display: flex;
    gap: 0.8rem;
    align-items: flex-start;
    padding: 1rem 1.1rem;
    border: 1px solid rgba(47, 106, 85, 0.26);
    border-radius: 14px;
    background: var(--green-soft);
}

.dc-clear-mark {
    display: grid;
    place-items: center;
    flex: 0 0 auto;
    width: 27px;
    height: 27px;
    border-radius: 50%;
    background: var(--green);
    color: #f7f1e6;
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.72rem;
}

.dc-clear h4 {
    margin: 0.05rem 0 0.2rem;
    font-size: 1.02rem;
}

.dc-clear p {
    margin: 0;
    color: var(--ink-soft);
    font-size: 0.8rem;
    line-height: 1.55;
}

.dc-tabs-space {
    height: 0.35rem;
}

/* Tabs */
.stTabs [data-baseweb="tab-list"] {
    gap: 0.3rem;
    border-bottom: 1px solid var(--line);
    background: transparent;
}

.stTabs [data-baseweb="tab"] {
    height: auto;
    padding: 0.72rem 0.95rem;
    border: 1px solid transparent;
    border-radius: 11px 11px 0 0;
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.64rem;
    letter-spacing: 0.02em;
    text-transform: uppercase;
}

.stTabs [data-baseweb="tab"]:hover {
    color: var(--ink);
    background: rgba(252, 248, 240, 0.65);
}

.stTabs [aria-selected="true"] {
    border-color: var(--line);
    border-bottom-color: var(--paper-high);
    background: var(--paper-high);
    color: var(--ink);
}

.stTabs [data-baseweb="tab-highlight"],
.stTabs [data-baseweb="tab-border"] {
    display: none;
}

.stTabs [data-baseweb="tab-panel"] {
    padding-top: 1.25rem;
}

/* Method */
.dc-method-grid {
    display: grid;
    grid-template-columns: 1.05fr 0.95fr;
    gap: 0.85rem;
}

.dc-method-card {
    padding: 1.2rem 1.3rem;
    border: 1px solid var(--line);
    border-radius: 15px;
    background: rgba(252, 248, 240, 0.7);
}

.dc-method-kicker {
    color: var(--accent);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.59rem;
    letter-spacing: 0.08em;
    text-transform: uppercase;
}

.dc-method-card h3 {
    margin: 0.6rem 0 0.65rem;
    font-size: 1.22rem;
    line-height: 1.12;
}

.dc-method-card ol,
.dc-method-card ul {
    margin: 0;
    padding-left: 1.1rem;
    color: var(--ink-soft);
    font-size: 0.82rem;
    line-height: 1.62;
}

.dc-method-card li {
    margin-bottom: 0.42rem;
}

.dc-method-card li:last-child {
    margin-bottom: 0;
}

.dc-run-settings {
    display: flex;
    flex-wrap: wrap;
    gap: 0.45rem;
    margin-top: 1rem;
}

.dc-run-settings span {
    display: inline-flex;
    gap: 0.4rem;
    align-items: baseline;
    max-width: 100%;
    padding: 0.35rem 0.55rem;
    border: 1px solid var(--line);
    border-radius: 8px;
    background: var(--paper-high);
    color: var(--ink);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.61rem;
    overflow-wrap: anywhere;
}

.dc-run-settings em {
    color: var(--ink-soft);
    font-style: normal;
    text-transform: uppercase;
}

/* Footer */
.dc-footer {
    display: flex;
    flex-wrap: wrap;
    justify-content: space-between;
    gap: 0.8rem;
    margin-top: 3.2rem;
    padding-top: 1rem;
    border-top: 1px solid var(--line-strong);
    color: var(--ink-soft);
    font-family: "IBM Plex Mono", Consolas, monospace;
    font-size: 0.6rem;
    letter-spacing: 0.04em;
    line-height: 1.5;
    text-transform: uppercase;
}

.dc-footer strong {
    color: var(--ink);
    font-weight: 500;
}

@keyframes dc-rise {
    from {
        opacity: 0;
        transform: translateY(14px);
    }
    to {
        opacity: 1;
        transform: translateY(0);
    }
}

@media (max-width: 1180px) {
    .dc-stats {
        grid-template-columns: repeat(3, minmax(0, 1fr));
    }
}

@media (max-width: 980px) {
    .dc-hero-grid,
    .dc-method-grid {
        grid-template-columns: 1fr;
    }

    .dc-art {
        max-width: 560px;
        transform: rotate(0.7deg);
    }

    .dc-stats {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }

    .dc-guide {
        grid-template-columns: 1fr;
    }

    .dc-guide-card {
        min-height: 0;
    }

    .dc-verdict {
        grid-template-columns: 1fr;
    }

    .dc-stamp {
        display: none;
    }
}

@media (max-width: 640px) {
    .block-container {
        padding-top: 1.8rem;
        padding-left: 1rem;
        padding-right: 1rem;
    }

    .dc-brandline {
        align-items: flex-start;
        flex-direction: column;
        gap: 0.35rem;
    }

    .dc-hero h1 {
        font-size: 3rem;
    }

    .dc-stats {
        grid-template-columns: 1fr;
    }

    .dc-priority-label {
        width: 100%;
    }

    .dc-section {
        align-items: flex-start;
        flex-wrap: wrap;
    }

    .dc-section-note {
        width: 100%;
        margin-left: 1.5rem;
    }
}

@media (prefers-reduced-motion: reduce) {
    *,
    *::before,
    *::after {
        animation: none !important;
        transition: none !important;
        scroll-behavior: auto !important;
    }
}
"""

EMPTY = "—"
RESULT_KEY = "dc_result"
TYPE_LABELS = {"number": "Число", "date": "Дата", "text": "Текст"}
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d.%m.%Y",
    "%d/%m/%Y",
    "%Y/%m/%d",
    "%d-%m-%Y",
    "%d.%m.%y",
    "%d/%m/%y",
    "%Y.%m.%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
)
CURRENCY_RE = re.compile(r"(?i)(?:₽|руб\.?|rub)")


@dataclass
class ComparisonResult:
    key: str
    fields: tuple[str, ...]
    field_kinds: dict[str, str]
    tolerance: Decimal
    relative_tolerance: Decimal
    trim: bool
    casefold: bool
    ignore_empty: bool
    ignore_date_format: bool
    file_a_name: str
    file_b_name: str
    rows_a: int
    rows_b: int
    diff: pd.DataFrame
    only_a: pd.DataFrame
    only_b: pd.DataFrame
    dup_a: pd.DataFrame
    dup_b: pd.DataFrame
    missing_a: pd.DataFrame
    missing_b: pd.DataFrame
    summary: pd.DataFrame
    report: bytes
    signature: str

    @property
    def empty_count(self) -> int:
        return len(self.missing_a) + len(self.missing_b)

    @property
    def dup_count(self) -> int:
        return len(self.dup_a) + len(self.dup_b)

    @property
    def dup_keys_a(self) -> int:
        return int(self.dup_a[self.key].nunique()) if not self.dup_a.empty else 0

    @property
    def dup_keys_b(self) -> int:
        return int(self.dup_b[self.key].nunique()) if not self.dup_b.empty else 0

    @property
    def only_count(self) -> int:
        return len(self.only_a) + len(self.only_b)

    @property
    def diff_count(self) -> int:
        return len(self.diff)

    @property
    def affected_keys(self) -> int:
        return int(self.diff["Ключ"].nunique()) if not self.diff.empty else 0

    @property
    def key_issues(self) -> int:
        return self.empty_count + self.dup_count

    @property
    def total_issues(self) -> int:
        return self.key_issues + self.only_count + self.diff_count

    @property
    def category_count(self) -> int:
        flags = (
            self.empty_count > 0,
            self.dup_count > 0,
            bool(len(self.only_a)),
            bool(len(self.only_b)),
            self.diff_count > 0,
        )
        return sum(flags)

    @property
    def tone(self) -> str:
        if self.total_issues == 0:
            return "clear"
        if self.key_issues:
            return "critical"
        return "warning"


def plural(count: int, one: str, few: str, many: str) -> str:
    number = abs(int(count))
    if number % 10 == 1 and number % 100 != 11:
        return one
    if 2 <= number % 10 <= 4 and not 12 <= number % 100 <= 14:
        return few
    return many


def categories_word(count: int) -> str:
    number = abs(int(count))
    return "категории" if number % 10 == 1 and number % 100 != 11 else "категориях"


def compact_html(markup: str) -> str:
    return " ".join(line.strip() for line in markup.splitlines() if line.strip())


def tolerance_text(tolerance: Decimal) -> str:
    return "точное совпадение" if tolerance == 0 else f"±{tolerance.normalize():f}"


def relative_tolerance_text(tolerance: Decimal) -> str:
    return f"{tolerance:.2f}%"


def inject_styles() -> None:
    st.markdown(f"<style>{CSS}</style>", unsafe_allow_html=True)


def section_head(index: str, title: str, note: str = "") -> None:
    note_html = f'<span class="dc-section-note">{html.escape(note)}</span>' if note else ""
    st.markdown(
        f"""
        <div class="dc-section">
            <span class="dc-section-index">{html.escape(index)}</span>
            <span class="dc-section-title">{html.escape(title)}</span>
            {note_html}
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_hero() -> None:
    st.markdown(
        """
        <header class="dc-hero">
            <div class="dc-brandline">
                <div class="dc-brand"><span class="dc-brand-mark"></span> DataCheck</div>
                <div>сверка реестров · прототип v0.1</div>
            </div>
            <div class="dc-hero-grid">
                <div class="dc-hero-copy">
                    <div class="dc-eyebrow">СВЕРКА CSV / DATA QUALITY <span class="dc-version">v0.1</span></div>
                    <h1>Сверка без<br><em>слепых зон.</em></h1>
                    <p class="dc-lede">
                        Сопоставьте два реестра по ключу и увидьте, где расходятся стоимость,
                        статусы, владельцы и сами записи. Без ручного перебора строк.
                    </p>
                    <div class="dc-meta">
                        <span>CSV / Excel</span>
                        <span>сводка + приоритеты</span>
                        <span>Excel-выгрузка</span>
                    </div>
                </div>
                <div class="dc-art" aria-hidden="true">
                    <div class="dc-art-top"><span>CHK / 2 FILES</span><span>RUN 001</span></div>
                    <div class="dc-filepair">
                        <div class="dc-mini-file">
                            <span class="dc-mini-tag">FILE / A</span>
                            <div class="dc-mini-row"><span>asset_id</span><strong>A-002</strong></div>
                            <div class="dc-mini-row"><span>cost</span><strong>8 500</strong></div>
                            <div class="dc-mini-row"><span>status</span><strong>active</strong></div>
                        </div>
                        <div class="dc-bridge">&ne;</div>
                        <div class="dc-mini-file dc-mini-file--b">
                            <span class="dc-mini-tag">FILE / B</span>
                            <div class="dc-mini-row"><span>asset_id</span><strong>A-002</strong></div>
                            <div class="dc-mini-row"><span>cost</span><strong class="dc-alert">8 900</strong></div>
                            <div class="dc-mini-row"><span>status</span><strong>active</strong></div>
                        </div>
                    </div>
                    <div class="dc-art-foot"><span>ключи · полнота · атрибуты</span><span>diff ready</span></div>
                </div>
            </div>
        </header>
        """,
        unsafe_allow_html=True,
    )


def render_upload_head(tag: str, title: str, copy: str) -> None:
    st.markdown(
        f"""
        <div class="dc-upload-head">
            <span class="dc-upload-tag">{html.escape(tag)}</span>
            <div>
                <div class="dc-upload-title">{html.escape(title)}</div>
                <p class="dc-upload-copy">{html.escape(copy)}</p>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_guide() -> None:
    st.markdown(
        """
        <section class="dc-guide">
            <article class="dc-guide-card">
                <span class="dc-guide-num">01</span>
                <h3>Загрузите два файла</h3>
                <p>CSV или Excel. Один файл — эталон, второй — реестр, который нужно с ним сверить.</p>
            </article>
            <article class="dc-guide-card">
                <span class="dc-guide-num">02</span>
                <h3>Выберите ключ</h3>
                <p>Это идентификатор записи. Для демо-набора подойдёт <code>asset_id</code>.</p>
            </article>
            <article class="dc-guide-card">
                <span class="dc-guide-num">03</span>
                <h3>Получите отчёт</h3>
                <p>Сводка покажет критичность и порядок разбора, а Excel-выгрузка сохранит детали.</p>
            </article>
        </section>
        <div class="dc-demo-hint"><strong>ДЕМО-НАБОР</strong> — файлы в папке <code>sample_data</code>, ключ <code>asset_id</code>. В них уже есть расхождения, дубль и пустые ID.</div>
        """,
        unsafe_allow_html=True,
    )


def render_footer() -> None:
    st.markdown(
        """
        <footer class="dc-footer">
            <span><strong>DataCheck</strong> / v0.1</span>
            <span>Python · Pandas · Streamlit</span>
            <span>Учебный прототип · синтетические данные</span>
        </footer>
        """,
        unsafe_allow_html=True,
    )


def load_file(uploaded) -> pd.DataFrame:
    if uploaded.name.lower().endswith(".csv"):
        raw = uploaded.getvalue()
        for encoding in ("utf-8-sig", "cp1251", "utf-8"):
            try:
                return pd.read_csv(
                    io.BytesIO(raw),
                    encoding=encoding,
                    sep=None,
                    engine="python",
                    dtype=str,
                    keep_default_na=False,
                )
            except Exception:
                continue
        raise ValueError("Не удалось прочитать CSV. Проверьте кодировку и разделитель.")
    try:
        return pd.read_excel(uploaded)
    except Exception as exc:
        raise ValueError(f"Не удалось прочитать Excel-файл: {exc}") from exc


def add_row_numbers(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    result["_row"] = range(2, len(result) + 2)
    return result


def show_value(value) -> str:
    if pd.isna(value):
        return EMPTY
    text = str(value).strip()
    return text if text else EMPTY


def is_blank(value) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return not str(value).strip()


def parse_number(value) -> Decimal | None:
    if is_blank(value):
        return None
    text = re.sub(r"\s+", "", str(value).replace("\u00a0", "").replace("\u202f", ""))
    text = CURRENCY_RE.sub("", text).replace("−", "-").replace("–", "-")
    if not text:
        return None

    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = text.lstrip("+")
    if not re.fullmatch(r"-?[\d.,]+", text) or not re.search(r"\d", text):
        return None

    if "," in text and "." in text:
        decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        integer_part, fraction = text.rsplit(decimal_separator, 1)
        integer_part = integer_part.replace(thousands_separator, "").replace(decimal_separator, "")
        text = f"{integer_part}.{fraction}"
    elif text.count(",") > 1:
        parts = text.split(",")
        text = "".join(parts) if all(len(part) == 3 for part in parts[1:]) else "".join(parts[:-1]) + "." + parts[-1]
    elif text.count(".") > 1:
        parts = text.split(".")
        text = "".join(parts) if all(len(part) == 3 for part in parts[1:]) else "".join(parts[:-1]) + "." + parts[-1]
    elif "," in text:
        text = text.replace(",", ".")

    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    if not number.is_finite():
        return None
    return -number if negative else number


def parse_date(value) -> date | None:
    if is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def detect_field_kind(values: list) -> str:
    non_empty = [value for value in values if not is_blank(value)]
    if not non_empty:
        return "text"
    if any(re.fullmatch(r"0\d+", str(value).strip()) for value in non_empty):
        return "text"
    if all(parse_number(value) is not None for value in non_empty):
        return "number"
    if all(parse_date(value) is not None for value in non_empty):
        return "date"
    return "text"


def numbers_within_tolerance(
    left: Decimal,
    right: Decimal,
    absolute_tolerance: Decimal,
    relative_tolerance_percent: Decimal,
) -> bool:
    difference = abs(left - right)
    if difference <= absolute_tolerance:
        return True
    denominator = max(abs(left), abs(right))
    if denominator == 0:
        return False
    return difference / denominator <= relative_tolerance_percent / Decimal("100")


def values_equal(
    left,
    right,
    kind: str,
    tolerance: Decimal,
    relative_tolerance: Decimal,
    trim: bool,
    casefold: bool,
    ignore_empty: bool,
    ignore_date_format: bool,
) -> bool:
    if is_blank(left) and is_blank(right):
        return True
    if is_blank(left) or is_blank(right):
        return ignore_empty
    if kind == "number":
        left_number, right_number = parse_number(left), parse_number(right)
        if left_number is not None and right_number is not None:
            return numbers_within_tolerance(left_number, right_number, tolerance, relative_tolerance)
    elif kind == "date" and ignore_date_format:
        left_date, right_date = parse_date(left), parse_date(right)
        if left_date is not None and right_date is not None:
            return left_date == right_date

    left_text, right_text = str(left), str(right)
    if trim:
        left_text, right_text = left_text.strip(), right_text.strip()
    if casefold:
        left_text, right_text = left_text.casefold(), right_text.casefold()
    return left_text == right_text


def make_signature(
    fa,
    fb,
    key: str,
    fields: list[str],
    tolerance: Decimal,
    relative_tolerance: Decimal,
    trim: bool,
    casefold: bool,
    ignore_empty: bool,
    ignore_date_format: bool,
) -> str:
    digest = hashlib.sha1()
    for uploaded in (fa, fb):
        digest.update(uploaded.name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(uploaded.getvalue())
        digest.update(b"\0")
    settings = repr(
        (
            key,
            tuple(fields),
            tolerance,
            relative_tolerance,
            trim,
            casefold,
            ignore_empty,
            ignore_date_format,
        )
    ).encode("utf-8")
    digest.update(settings)
    return digest.hexdigest()


def public_table(frame: pd.DataFrame) -> pd.DataFrame:
    columns = [column for column in ("_row", *frame.columns) if column in frame.columns]
    result = frame.loc[:, list(dict.fromkeys(columns))].copy()
    return result.rename(columns={"_row": "Строка"})


def with_file_column(frame: pd.DataFrame, label: str) -> pd.DataFrame:
    result = public_table(frame)
    result.insert(0, "Файл", label)
    return result


def fill_empty_keys(frame: pd.DataFrame, key: str) -> pd.DataFrame:
    result = frame.copy()
    if key in result.columns:
        result[key] = result[key].astype("string").fillna(EMPTY).replace("", EMPTY)
    return result


def empty_keys_table(result: ComparisonResult) -> pd.DataFrame:
    frames = []
    for label, frame in (("A", result.missing_a), ("B", result.missing_b)):
        if not frame.empty:
            frames.append(fill_empty_keys(with_file_column(frame, label), result.key))
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True)


def duplicate_rows_table(result: ComparisonResult) -> pd.DataFrame:
    frames = []
    for label, frame in (("A", result.dup_a), ("B", result.dup_b)):
        if not frame.empty:
            frames.append(with_file_column(frame, label))
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True).sort_values(["Файл", result.key, "Строка"])


def dataframe_height(row_count: int, maximum: int = 430) -> int:
    return min(maximum, max(148, 38 + 35 * max(row_count, 1)))


def build_excel(
    summary: pd.DataFrame,
    parameters: pd.DataFrame,
    diff: pd.DataFrame,
    only_a: pd.DataFrame,
    only_b: pd.DataFrame,
    dup_a: pd.DataFrame,
    dup_b: pd.DataFrame,
    missing_a: pd.DataFrame,
    missing_b: pd.DataFrame,
    key: str,
) -> bytes:
    sheets = {
        "Сводка": summary,
        "Параметры": parameters,
        "Расхождения": diff,
        "Только_A": public_table(only_a),
        "Только_B": public_table(only_b),
        "Дубли_A": public_table(dup_a),
        "Дубли_B": public_table(dup_b),
        "Без_ключа_A": fill_empty_keys(public_table(missing_a), key),
        "Без_ключа_B": fill_empty_keys(public_table(missing_b), key),
    }
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        header_fill = PatternFill("solid", fgColor="14221D")
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_properties.tabColor = "C34A2C"
            worksheet.row_dimensions[1].height = 22
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(name="Calibri", bold=True, color="F7F1E4", size=10)
                cell.alignment = Alignment(vertical="center")
            for index, column in enumerate(worksheet.columns, start=1):
                lengths = [len(str(cell.value)) for cell in column if cell.value is not None]
                width = min(max(max(lengths, default=10) + 2, 11), 48)
                worksheet.column_dimensions[get_column_letter(index)].width = width
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top")
    return output.getvalue()


def compare_frames(
    a: pd.DataFrame,
    b: pd.DataFrame,
    key: str,
    compare_fields: list[str],
    tolerance: Decimal,
    relative_tolerance: Decimal,
    trim: bool,
    casefold: bool,
    ignore_empty: bool,
    ignore_date_format: bool,
    file_a_name: str,
    file_b_name: str,
    signature: str,
) -> ComparisonResult:
    a = add_row_numbers(a)
    b = add_row_numbers(b)
    compared = [field for field in compare_fields if field != key and field in a.columns and field in b.columns]

    for frame in (a, b):
        frame[key] = frame[key].astype("string").str.strip()

    blank_a = a[key].isna() | (a[key] == "")
    blank_b = b[key].isna() | (b[key] == "")
    missing_a = a[blank_a].copy()
    missing_b = b[blank_b].copy()
    valid_a = a[~blank_a].copy()
    valid_b = b[~blank_b].copy()

    dup_a = valid_a[valid_a.duplicated(key, keep=False)].sort_values([key, "_row"])
    dup_b = valid_b[valid_b.duplicated(key, keep=False)].sort_values([key, "_row"])
    unique_a = valid_a.drop_duplicates(key, keep="first").set_index(key, drop=False)
    unique_b = valid_b.drop_duplicates(key, keep="first").set_index(key, drop=False)

    keys_a = set(unique_a.index.astype(str))
    keys_b = set(unique_b.index.astype(str))
    only_keys_a = sorted(keys_a - keys_b)
    only_keys_b = sorted(keys_b - keys_a)
    only_a = unique_a.loc[only_keys_a].reset_index(drop=True) if only_keys_a else unique_a.iloc[0:0].reset_index(drop=True)
    only_b = unique_b.loc[only_keys_b].reset_index(drop=True) if only_keys_b else unique_b.iloc[0:0].reset_index(drop=True)

    common_keys = sorted(keys_a & keys_b)
    field_kinds = {}
    for field in compared:
        samples = []
        for common_key in common_keys:
            samples.append(unique_a.at[common_key, field])
            samples.append(unique_b.at[common_key, field])
        field_kinds[field] = detect_field_kind(samples)

    changes = []
    for common_key in common_keys:
        row_a = unique_a.loc[common_key]
        row_b = unique_b.loc[common_key]
        for field in compared:
            kind = field_kinds[field]
            if not values_equal(
                row_a[field],
                row_b[field],
                kind,
                tolerance,
                relative_tolerance,
                trim,
                casefold,
                ignore_empty,
                ignore_date_format,
            ):
                changes.append(
                    {
                        "Ключ": common_key,
                        "Поле": field,
                        "Тип": TYPE_LABELS[kind],
                        "Файл A": show_value(row_a[field]),
                        "Файл B": show_value(row_b[field]),
                        "Строка A": int(row_a["_row"]),
                        "Строка B": int(row_b["_row"]),
                    }
                )

    diff = pd.DataFrame(
        changes,
        columns=["Ключ", "Поле", "Тип", "Файл A", "Файл B", "Строка A", "Строка B"],
    )
    for column in ("Ключ", "Поле", "Тип", "Файл A", "Файл B"):
        diff[column] = diff[column].astype("string")
    diff["Строка A"] = diff["Строка A"].astype("Int64")
    diff["Строка B"] = diff["Строка B"].astype("Int64")

    summary = pd.DataFrame(
        [
            ["Строки в A", len(a)],
            ["Строки в B", len(b)],
            ["Только в A", len(only_a)],
            ["Только в B", len(only_b)],
            ["Ключей с дублями в A", dup_a[key].nunique()],
            ["Ключей с дублями в B", dup_b[key].nunique()],
            ["Строк без ключа в A", len(missing_a)],
            ["Строк без ключа в B", len(missing_b)],
            ["Расхождений полей", len(diff)],
            ["Затронуто ключей", len(diff["Ключ"].dropna().unique())],
        ],
        columns=["Проверка", "Количество"],
    )
    field_types = "; ".join(
        f"{field} — {TYPE_LABELS[field_kinds[field]]}" for field in compared
    ) or EMPTY
    parameters = pd.DataFrame(
        [
            ["Файл A", file_a_name],
            ["Файл B", file_b_name],
            ["Ключ", key],
            ["Поля сравнения", ", ".join(compared) if compared else EMPTY],
            ["Типы полей", field_types],
            ["Допуск для чисел", tolerance_text(tolerance)],
            ["Относительный допуск", relative_tolerance_text(relative_tolerance)],
            ["Пробелы по краям текста", "игнорируются" if trim else "учитываются"],
            ["Регистр текста", "игнорируется" if casefold else "учитывается"],
            ["Пустые значения", "игнорируются" if ignore_empty else "учитываются"],
            ["Формат дат", "игнорируется" if ignore_date_format else "учитывается"],
        ],
        columns=["Параметр", "Значение"],
    )

    report = build_excel(
        summary,
        parameters,
        diff,
        only_a,
        only_b,
        dup_a,
        dup_b,
        missing_a,
        missing_b,
        key,
    )

    return ComparisonResult(
        key=key,
        fields=tuple(compared),
        field_kinds=field_kinds,
        tolerance=tolerance,
        relative_tolerance=relative_tolerance,
        trim=trim,
        casefold=casefold,
        ignore_empty=ignore_empty,
        ignore_date_format=ignore_date_format,
        file_a_name=file_a_name,
        file_b_name=file_b_name,
        rows_a=len(a),
        rows_b=len(b),
        diff=diff,
        only_a=only_a,
        only_b=only_b,
        dup_a=dup_a,
        dup_b=dup_b,
        missing_a=missing_a,
        missing_b=missing_b,
        summary=summary,
        report=report,
        signature=signature,
    )


def render_runline(result: ComparisonResult) -> None:
    file_a = html.escape(result.file_a_name)
    file_b = html.escape(result.file_b_name)
    key = html.escape(result.key)
    st.markdown(
        f"""
        <div class="dc-runline">
            <span class="dc-runline-label">Сверка</span>
            <span class="dc-runline-file" title="{file_a}">{file_a}</span>
            <span class="dc-runline-x">×</span>
            <span class="dc-runline-file" title="{file_b}">{file_b}</span>
            <span class="dc-runline-divider"></span>
            <span>ключ</span><strong>{key}</strong>
            <span class="dc-runline-divider"></span>
            <span>{len(result.fields)} {plural(len(result.fields), "поле", "поля", "полей")}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_verdict(result: ComparisonResult) -> None:
    if result.tone == "clear":
        title = "Файлы согласованы."
        body = "Проверка ключей, состава записей и значений не нашла замечаний."
        stamp = "CLEAR"
        stamp_note = "NO ISSUES"
    else:
        total = result.total_issues
        category_count = result.category_count
        title = (
            f"Найдено {total} {plural(total, 'замечание', 'замечания', 'замечаний')}<br>"
            f"в {category_count} {categories_word(category_count)}."
        )
        if result.key_issues:
            body = "Сначала разберите ключи: пустые идентификаторы и дубли мешают однозначному сопоставлению."
        elif result.only_count:
            body = "Сверьте записи, которых нет в одном из файлов, и только затем переходите к значениям."
        else:
            body = "Ключи и состав записей в порядке. Осталось разобрать отличия в значениях."
        stamp = "REVIEW"
        stamp_note = "REQUIRED"

    st.markdown(
        f"""
        <section class="dc-verdict dc-verdict--{result.tone}">
            <div>
                <div class="dc-verdict-kicker">Проверка завершена</div>
                <h2>{title}</h2>
                <p>{html.escape(body)}</p>
            </div>
            <div class="dc-stamp" aria-hidden="true">
                <span>{stamp_note}</span>
                <strong>{stamp}</strong>
                <span>datacheck / 0.1</span>
            </div>
        </section>
        """,
        unsafe_allow_html=True,
    )


def stat_card(index: str, label: str, value: str, foot: str, tone: str) -> str:
    return f"""
    <article class="dc-stat dc-stat--{tone}">
        <div>
            <div class="dc-stat-index">{index}</div>
            <div class="dc-stat-label">{html.escape(label)}</div>
            <div class="dc-stat-value">{value}</div>
        </div>
        <div class="dc-stat-foot">{html.escape(foot)}</div>
    </article>
    """


def render_stats(result: ComparisonResult) -> None:
    duplicate_value = result.dup_keys_a + result.dup_keys_b
    rows_value = f'{result.rows_a}<span class="dc-slash">/</span>{result.rows_b}'
    cards = [
        stat_card("01", "Строки", rows_value, "A / B", "neutral"),
        stat_card(
            "02",
            "Только в одном",
            str(result.only_count),
            f"A {len(result.only_a)} · B {len(result.only_b)}",
            "warning" if result.only_count else "clear",
        ),
        stat_card(
            "03",
            "Дубли ключей",
            str(duplicate_value),
            f"A {result.dup_keys_a} · B {result.dup_keys_b} · {result.dup_count} {plural(result.dup_count, 'строка', 'строки', 'строк')}",
            "critical" if result.dup_count else "clear",
        ),
        stat_card(
            "04",
            "Пустые ID",
            str(result.empty_count),
            f"A {len(result.missing_a)} · B {len(result.missing_b)}",
            "critical" if result.empty_count else "clear",
        ),
        stat_card(
            "05",
            "Расхождения",
            str(result.diff_count),
            f"{result.affected_keys} {plural(result.affected_keys, 'ключ', 'ключа', 'ключей')} затронуто",
            "warning" if result.diff_count else "clear",
        ),
    ]
    st.markdown(compact_html(f'<div class="dc-stats">{"".join(cards)}</div>'), unsafe_allow_html=True)


def render_priority(result: ComparisonResult) -> None:
    key_tone = "critical" if result.key_issues else "clear"
    only_tone = "warning" if result.only_count else "clear"
    diff_tone = "warning" if result.diff_count else "clear"
    st.markdown(
        f"""
        <div class="dc-priority">
            <span class="dc-priority-label">Порядок разбора</span>
            <span class="dc-chip dc-chip--{key_tone}"><i>01</i> Ключи <b>{result.key_issues}</b></span>
            <span class="dc-chip dc-chip--{only_tone}"><i>02</i> Полнота <b>{result.only_count}</b></span>
            <span class="dc-chip dc-chip--{diff_tone}"><i>03</i> Значения <b>{result.diff_count}</b></span>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_clear_state(title: str, text: str) -> None:
    st.markdown(
        f"""
        <div class="dc-clear">
            <span class="dc-clear-mark">✓</span>
            <div>
                <h4>{html.escape(title)}</h4>
                <p>{html.escape(text)}</p>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_table_head(title: str, meta: str) -> None:
    st.markdown(
        f"""
        <div class="dc-table-head">
            <span class="dc-table-title">{html.escape(title)}</span>
            <span class="dc-table-meta">{html.escape(meta)}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )


def display_frame(frame: pd.DataFrame, height: int | None = None) -> None:
    public = public_table(frame)
    column_config = {}
    if "Строка" in public.columns:
        column_config["Строка"] = st.column_config.NumberColumn("Строка", width="small")
    st.dataframe(
        public,
        width="stretch",
        hide_index=True,
        height=height or dataframe_height(len(public)),
        column_config=column_config,
    )


def render_diff_tab(result: ComparisonResult) -> None:
    st.markdown(
        '<p class="dc-tab-lede">Только поля с разными значениями после выбранных правил сравнения. Номер строки помогает найти запись в исходном файле.</p>',
        unsafe_allow_html=True,
    )
    if result.diff.empty:
        render_clear_state("Расхождений значений нет", "Все выбранные поля в общих записях совпали.")
        return
    st.dataframe(
        result.diff,
        width="stretch",
        hide_index=True,
        height=dataframe_height(len(result.diff)),
        column_config={
            "Ключ": st.column_config.TextColumn("Ключ", width="small"),
            "Поле": st.column_config.TextColumn("Поле", width="small"),
            "Тип": st.column_config.TextColumn("Тип", width="small"),
            "Файл A": st.column_config.TextColumn("Файл A", width="medium"),
            "Файл B": st.column_config.TextColumn("Файл B", width="medium"),
            "Строка A": st.column_config.NumberColumn("Строка A", width="small"),
            "Строка B": st.column_config.NumberColumn("Строка B", width="small"),
        },
    )
    st.markdown(
        '<div class="dc-table-note">Тип поля определён автоматически по общим записям. Строка 1 — заголовок файла, поэтому данные начинаются со строки 2. Файлы не изменяются.</div>',
        unsafe_allow_html=True,
    )


def render_only_tab(result: ComparisonResult) -> None:
    st.markdown(
        '<p class="dc-tab-lede">Записи, которые есть только в одном файле. Сначала проверьте, потерялась ли строка при выгрузке или она появилась позже.</p>',
        unsafe_allow_html=True,
    )
    left, right = st.columns(2, gap="large")
    with left:
        render_table_head("Только в A", f"{len(result.only_a)} {plural(len(result.only_a), 'запись', 'записи', 'записей')}")
        if result.only_a.empty:
            render_clear_state("Лишних записей нет", "Все ключи из A нашлись в B.")
        else:
            display_frame(result.only_a)
    with right:
        render_table_head("Только в B", f"{len(result.only_b)} {plural(len(result.only_b), 'запись', 'записи', 'записей')}")
        if result.only_b.empty:
            render_clear_state("Лишних записей нет", "Все ключи из B нашлись в A.")
        else:
            display_frame(result.only_b)


def render_keys_tab(result: ComparisonResult) -> None:
    st.markdown(
        '<p class="dc-tab-lede">Ключи проверяются первыми: пустой или повторяющийся идентификатор делает сопоставление строк неоднозначным.</p>',
        unsafe_allow_html=True,
    )
    render_table_head(
        "Пустые идентификаторы",
        f"{result.empty_count} {plural(result.empty_count, 'строка', 'строки', 'строк')}",
    )
    if result.empty_count == 0:
        render_clear_state("Пустых ключей нет", "В обеих выгрузках у каждой строки есть идентификатор.")
    else:
        display_frame(empty_keys_table(result))
        st.markdown(
            f'<div class="dc-table-note">Строки без значения в поле <b>{html.escape(result.key)}</b> исключены из сравнения значений. Их нужно исправить в источнике.</div>',
            unsafe_allow_html=True,
        )

    st.write("")
    render_table_head(
        "Дубликаты ключей",
        f"{result.dup_count} {plural(result.dup_count, 'строка', 'строки', 'строк')}",
    )
    if result.dup_count == 0:
        render_clear_state("Дублей ключей нет", "Каждый идентификатор встречается не более одного раза.")
    else:
        display_frame(duplicate_rows_table(result))
        st.markdown(
            '<div class="dc-table-note">В v0.1 для сравнения значений такого ключа используется первая строка. Отнесите результат по этим ключам к ручной проверке.</div>',
            unsafe_allow_html=True,
        )


def render_method_tab(result: ComparisonResult) -> None:
    fields = ", ".join(html.escape(field) for field in result.fields) or EMPTY
    types = " · ".join(
        f"{html.escape(field)} — {TYPE_LABELS.get(kind, 'Текст').lower()}"
        for field, kind in result.field_kinds.items()
    ) or EMPTY
    tolerance = tolerance_text(result.tolerance)
    relative_tolerance = relative_tolerance_text(result.relative_tolerance)
    spaces = "игнорируются" if result.trim else "учитываются"
    case = "игнорируется" if result.casefold else "учитывается"
    empty_rule = "игнорируются" if result.ignore_empty else "учитываются"
    date_rule = "игнорируется" if result.ignore_date_format else "учитывается"
    date_note = (
        "Формат дат игнорируется: сравнивается календарный день без времени"
        if result.ignore_date_format
        else "Формат дат учитывается: даты сравниваются как текст, включая время"
    )
    st.markdown(
        f"""
        <div class="dc-method-grid">
            <article class="dc-method-card">
                <span class="dc-method-kicker">Логика проверки</span>
                <h3>Сначала ключи, затем записи, потом значения</h3>
                <ol>
                    <li>Ключи очищаются от пробелов по краям, пустые идентификаторы выносятся отдельно.</li>
                    <li>Записи сопоставляются по выбранному ключу; для повторяющегося ключа в сравнении значений используется первая строка.</li>
                    <li>Тип поля определяется автоматически: числа сравниваются по значению с абсолютным или относительным допуском, даты — по выбранному правилу формата, текст — по правилам ниже.</li>
                    <li>Записи только в одном файле и расхождения значений попадают в отдельные разделы.</li>
                </ol>
            </article>
            <article class="dc-method-card">
                <span class="dc-method-kicker">Границы v0.1</span>
                <h3>Что стоит знать перед выводами</h3>
                <ul>
                    <li>Число считается равным, если проходит хотя бы один допуск: абсолютный (в единицах поля) или относительный (в % от большего по модулю значения).</li>
                    <li>При нулевых допусках требуется точное совпадение. {date_note}.</li>
                    <li>Для текста пробелы по краям {spaces}, регистр {case}. Смешанные типы сравниваются как текст; результат требует ручной проверки.</li>
                    <li>Пустые значения {empty_rule}. Пустые ключи всегда выносятся в отдельную проверку и правилом не скрываются.</li>
                    <li>Excel-отчёт содержит сводку, параметры запуска и проблемные строки.</li>
                </ul>
            </article>
        </div>
        <div class="dc-run-settings">
            <span><em>Ключ</em> <b>{html.escape(result.key)}</b></span>
            <span><em>Поля</em> <b>{fields}</b></span>
            <span><em>Типы</em> <b>{types}</b></span>
            <span><em>Допуск</em> <b>{tolerance}</b></span>
            <span><em>Отн. допуск</em> <b>{relative_tolerance}</b></span>
            <span><em>Пробелы</em> <b>{spaces}</b></span>
            <span><em>Регистр</em> <b>{case}</b></span>
            <span><em>Пустые</em> <b>{empty_rule}</b></span>
            <span><em>Формат дат</em> <b>{date_rule}</b></span>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_results(result: ComparisonResult) -> None:
    section_head("03", "Отчёт", "сводка · приоритеты · выгрузка")
    render_runline(result)
    render_verdict(result)
    render_stats(result)
    render_priority(result)

    action_left, action_right = st.columns([2.1, 1], gap="large")
    with action_left:
        st.markdown(
            '<div class="dc-action-note">Excel-отчёт · сводка, параметры запуска и все проблемные строки.<br>Исходные файлы не изменяются.</div>',
            unsafe_allow_html=True,
        )
    with action_right:
        st.download_button(
            "Скачать отчёт · XLSX",
            data=result.report,
            file_name="datacheck_report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
            key="dc_download_report",
        )

    st.markdown('<div class="dc-tabs-space"></div>', unsafe_allow_html=True)
    tabs = st.tabs(
        [
            f"01 · Расхождения · {result.diff_count}",
            f"02 · Полнота · {result.only_count}",
            f"03 · Ключи · {result.key_issues}",
            "04 · Методика",
        ]
    )
    with tabs[0]:
        render_diff_tab(result)
    with tabs[1]:
        render_only_tab(result)
    with tabs[2]:
        render_keys_tab(result)
    with tabs[3]:
        render_method_tab(result)


def main() -> None:
    inject_styles()
    render_hero()

    section_head("01", "Исходные файлы", "CSV · XLSX · два файла")
    left, right = st.columns(2, gap="large")
    with left:
        render_upload_head(
            "A",
            "Операционный файл",
            "Реестр, который вы считаете эталонным: например, операционный учёт лицензий.",
        )
        file_a = st.file_uploader(
            "Файл A",
            type=["csv", "xlsx", "xls"],
            label_visibility="collapsed",
            key="dc_file_a",
        )
    with right:
        render_upload_head(
            "B",
            "Файл для сверки",
            "Вторая выгрузка: бухгалтерия, закупки или другой источник, который нужно сравнить.",
        )
        file_b = st.file_uploader(
            "Файл B",
            type=["csv", "xlsx", "xls"],
            label_visibility="collapsed",
            key="dc_file_b",
        )

    if not (file_a and file_b):
        render_guide()
        render_footer()
        return

    try:
        a = load_file(file_a)
        b = load_file(file_b)
    except Exception as exc:
        st.error(f"Ошибка чтения файла: {exc}")
        render_footer()
        return

    preview_left, preview_right = st.columns(2, gap="large")
    for column, label, uploaded, frame in (
        (preview_left, "A", file_a, a),
        (preview_right, "B", file_b, b),
    ):
        with column:
            with st.expander(f"Предпросмотр {label} · {uploaded.name}"):
                st.caption(f"{len(frame)} строк · {len(frame.columns)} столбцов")
                st.dataframe(frame.head(8), width="stretch", hide_index=True, height=280)

    common = [column for column in a.columns if column in b.columns]
    if not common:
        st.error("У файлов нет одинаковых названий столбцов. Выберите файлы с общей схемой.")
        render_footer()
        return

    section_head("02", "Параметры аудита", "ключ · допуски · правила сравнения")
    with st.container(border=True):
        key_column, fields_column = st.columns([1, 1.6], gap="large")
        with key_column:
            key = st.selectbox(
                "Ключ сопоставления",
                common,
                help="Столбец, который однозначно идентифицирует запись. В демо-наборе это asset_id.",
            )
        with fields_column:
            available_fields = [column for column in common if column != key]
            compare_fields = st.multiselect(
                "Поля для сравнения",
                available_fields,
                default=available_fields,
                help="Оставьте только те атрибуты, расхождения в которых важны для проверки.",
                wrap=True,
            )

        absolute_column, relative_column, tolerance_note_column = st.columns(
            [1, 1, 2],
            gap="large",
            vertical_alignment="bottom",
        )
        with absolute_column:
            tolerance_value = st.number_input(
                "Допуск для чисел",
                min_value=0.0,
                value=0.01,
                step=0.01,
                format="%.2f",
                help="Максимальная абсолютная разница, которая не считается расхождением. 0 — только точное совпадение.",
            )
        with relative_column:
            relative_tolerance_value = st.number_input(
                "Относительный допуск, %",
                min_value=0.0,
                value=0.0,
                step=0.01,
                format="%.2f",
                help="Максимальная относительная разница в процентах от большего по модулю значения. 0 отключает правило.",
            )
        with tolerance_note_column:
            st.markdown(
                '<div class="dc-tolerance-note">Значения совпадут, если проходит хотя бы один допуск: абсолютный или относительный.</div>',
                unsafe_allow_html=True,
            )

        option_spacing, option_case = st.columns(2, gap="large")
        with option_spacing:
            trim = st.checkbox(
                "Игнорировать пробелы по краям текста",
                value=True,
                help="Пробелы в начале и конце не учитываются только у текстовых полей. Ключи всегда очищаются перед сопоставлением.",
            )
        with option_case:
            casefold = st.checkbox(
                "Игнорировать регистр текстовых полей",
                value=True,
                help="Регистр не учитывается только для полей, распознанных как текст. Числа и даты сравниваются по своим правилам.",
            )
        option_date, option_empty = st.columns(2, gap="large")
        with option_date:
            ignore_date_format = st.checkbox(
                "Игнорировать различия в формате дат",
                value=True,
                help="Сравнивать даты по календарному дню. Например, 2024-04-12 и 12.04.2024 совпадут. Если выключить, даты сравниваются как текст.",
            )
        with option_empty:
            ignore_empty = st.checkbox(
                "Игнорировать пустые значения в полях",
                value=False,
                help="Не считать расхождением пару «пусто — значение». Пустые ключи проверяются отдельно и этим правилом не скрываются.",
            )

        tolerance = Decimal(str(tolerance_value if tolerance_value is not None else 0))
        relative_tolerance = Decimal(
            str(relative_tolerance_value if relative_tolerance_value is not None else 0)
        )

        button_column, note_column = st.columns([1, 2], gap="large")
        with button_column:
            compare_clicked = st.button("Сверить файлы →", type="primary", width="stretch")
        with note_column:
            st.markdown(
                '<div class="dc-inline-note">Допуски применяются только к распознанным числам. Файлы остаются на месте.</div>',
                unsafe_allow_html=True,
            )

    signature = make_signature(
        file_a,
        file_b,
        key,
        compare_fields,
        tolerance,
        relative_tolerance,
        trim,
        casefold,
        ignore_empty,
        ignore_date_format,
    )
    stored_result = st.session_state.get(RESULT_KEY)

    if compare_clicked:
        try:
            with st.spinner("Сопоставляю записи и собираю отчёт…"):
                stored_result = compare_frames(
                    a,
                    b,
                    key,
                    compare_fields,
                    tolerance,
                    relative_tolerance,
                    trim,
                    casefold,
                    ignore_empty,
                    ignore_date_format,
                    file_a.name,
                    file_b.name,
                    signature,
                )
            st.session_state[RESULT_KEY] = stored_result
        except Exception as exc:
            st.error(f"Не удалось выполнить сверку: {exc}")
            render_footer()
            return

    if stored_result is not None:
        if stored_result.signature != signature:
            st.markdown(
                '<div class="dc-change-note">Параметры или файлы изменились. Запустите сверку ещё раз, чтобы обновить отчёт.</div>',
                unsafe_allow_html=True,
            )
            render_footer()
            return
        render_results(stored_result)
    else:
        render_clear_state(
            "Файлы готовы к проверке",
            "Выберите ключ и поля, затем нажмите «Сверить файлы».",
        )

    render_footer()


main()
